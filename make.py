"""Render the Hang email as a .docx.

Local convenience only -- python-docx is not used anywhere in the benchmark and
is not required on Argon.

The source of truth stays `docs/email_to_hang.md`; this reads it and produces a
Word file so the table arrives as a real table rather than pipe characters. The
"Notes for you, not for the email" section is dropped, since the document is
meant to be sendable as-is.

Usage:
    .venv\\Scripts\\python.exe make_email_docx.py
"""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "docs" / "email_to_hang.md"
OUTPUT = ROOT / "docs" / "Stage1_results_Hang.docx"

DROP_AFTER = "Notes for you, not for the email"


def add_runs(paragraph, text: str) -> None:
    """Emit text, honouring **bold**, *italic* and `code` spans."""

    for piece in re.split(r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`)", text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            paragraph.add_run(piece[1:-1]).italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        else:
            paragraph.add_run(piece)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main() -> int:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()

    # Trim the private notes and the draft's own front matter.
    end = next((i for i, l in enumerate(lines) if DROP_AFTER in l), len(lines))
    lines = lines[:end]
    start = next((i for i, l in enumerate(lines) if l.startswith("**Subject:**")), 0)
    lines = lines[start:]

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    index = 0
    while index < len(lines):
        line = lines[index].rstrip()

        if not line or set(line) <= {"-"} and len(line) > 2:
            index += 1
            continue

        # Table: a header row followed by a separator of dashes and pipes.
        if (line.startswith("|") and index + 1 < len(lines)
                and re.fullmatch(r"\|[\s:\-|]+\|", lines[index + 1].strip())):
            header = split_row(line)
            index += 2
            body: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                body.append(split_row(lines[index]))
                index += 1

            table = document.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, header):
                cell.text = ""
                add_runs(cell.paragraphs[0], f"**{text}**")
            for record in body:
                cells = table.add_row().cells
                for cell, text in zip(cells, record):
                    cell.text = ""
                    add_runs(cell.paragraphs[0], text)
                    cell.paragraphs[0].runs and setattr(
                        cell.paragraphs[0].runs[0].font, "size", Pt(9.5))
            document.add_paragraph()
            continue

        # Absorb wrapped continuation lines FIRST, so a list item keeps its own
        # continuation instead of spilling into a separate paragraph. A new block
        # starts at a heading, a bullet, a table row, a numbered item, or a blank.
        buffer = [line]
        while index + 1 < len(lines):
            nxt = lines[index + 1]
            if (not nxt.strip() or nxt.startswith(("#", "|"))
                    or re.match(r"^\s*[-*]\s", nxt) or re.match(r"^\s*\d+\.\s", nxt)):
                break
            index += 1
            buffer.append(nxt.strip())
        block = " ".join(buffer)

        if block.startswith("## "):
            document.add_heading(block[3:].strip(), level=2)
        elif block.startswith("# "):
            document.add_heading(block[2:].strip(), level=1)
        elif re.match(r"^\d+\.\s", block):
            paragraph = document.add_paragraph(style="List Number")
            add_runs(paragraph, re.sub(r"^\d+\.\s", "", block))
        elif re.match(r"^[-*]\s", block):
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, block[2:])
        elif block.startswith("**Subject:**"):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(paragraph, block)
            document.add_paragraph()
        else:
            add_runs(document.add_paragraph(), block)

        index += 1

    document.save(OUTPUT)
    print(f"Saved {OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
